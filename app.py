import os
import io
import pandas as pd
import requests
import openpyxl
import streamlit as st
from roboflow import Roboflow
from openpyxl.styles import PatternFill, Font

# ==========================================
# 1. СЛОВАРЬ КРАСИВЫХ НАЗВАНИЙ ПРЕПАРАТОВ
# ==========================================
PRODUCT_NAMES = {
    "cardiom_150_100": "Кардиомагнил 150мг",
    "ent_2_10": "Энтерожермина 2мл №10",
    "ent_2_12": "Энтерожермина 2мл №12",
    "ent_4_10": "Энтерожермина 4мл №10",
    "magneb6_60": "Магне В6 №60",
    "nospa_40_24": "Но-шпа 40мг №24",
    "snup_01_15": "Снуп 0.1% 15мл",
    "zodac_10_30": "Зодак 10мг №30",
    "las_15_5_100": "Лазолван 15мг",
    "snup_005_15": "Снуп 0.05% 15мл",
    "festal_20": "Фестал №20",
    "ent_6_9": "Энтерожермина 6мл №9"
}

st.set_page_config(page_title="Аудит Выкладки", layout="centered")

st.image("logo.png", width=250)
st.title("💊 Автоматический аудит выкладки")
st.write("Загрузите файл матрицы (план) и файл с фотоотчетом (ссылками), чтобы нейросеть проверила наличие препаратов.")

st.subheader("1. Загрузка данных")
matrix_file = st.file_uploader("Загрузите матрицу (matrix.xlsx)", type=["xlsx"])
report_file = st.file_uploader("Загрузите фотоотчет (Oson...xlsm или xlsx)", type=["xlsx", "xlsm"])

api_key_input = st.text_input("Введите ваш API-ключ Roboflow", type="password")

if st.button("🚀 Запустить проверку", type="primary"):
    if not matrix_file or not report_file or not api_key_input:
        st.error("Пожалуйста, загрузите оба файла и введите API-ключ!")
    else:
        try:
            # st.info("Подключение к нейросети...")
            # rf = Roboflow(api_key=api_key_input)
            # project = rf.workspace().project("uz_ir_pharmacy") # Вернули автоматический поиск workspace
            # model = project.version(4) # ВНИМАНИЕ: Убрали .model в конце!
            
            # # Ставим защиту: если Roboflow вдруг отдаст пустоту, скрипт сразу скажет почему
            # if model is None:
            #     st.error("Критическая ошибка: Roboflow не отдал модель. Проверь, что 4 версия точно обучена!")
            #     st.stop()
            
            st.info("Чтение планов матрицы...")
            df_matrix = pd.read_excel(matrix_file)
            df_matrix = df_matrix.dropna(subset=['roboflow_name'])

            matrix_plans = {}
            all_unique_products = set()
            
            # Сборка планов: Месяц -> Полка -> Вариант -> Продукты
            for index, row in df_matrix.iterrows():
                try:
                    month = int(row['Month'])
                except:
                    month = 1 
                
                shelf_name = str(row['Column Name']).strip()
                
                try:
                    option = str(int(row['Option'])).strip()
                except:
                    option = str(row['Option']).strip() if pd.notna(row['Option']) else "1"
                    
                product = str(row['roboflow_name']).strip()
                target = int(row['Target packs'])
                
                if month not in matrix_plans:
                    matrix_plans[month] = {}
                if shelf_name not in matrix_plans[month]:
                    matrix_plans[month][shelf_name] = {}
                if option not in matrix_plans[month][shelf_name]:
                    matrix_plans[month][shelf_name][option] = {}
                    
                matrix_plans[month][shelf_name][option][product] = target
                all_unique_products.add(product)

            all_unique_products = sorted(list(all_unique_products))

            st.info("Открытие фотоотчета...")
            wb = openpyxl.load_workbook(report_file, data_only=True)
            ws = wb.active
            
            max_rows = ws.max_row + 1 
            master_results = []
            
            progress_bar = st.progress(0)
            status_text = st.empty()
            error_log = st.empty() # Специальный блок для вывода скрытых ошибок

            total_items = max_rows - 2
            if total_items <= 0:
                st.warning("В загруженном отчете нет данных для проверки.")
                st.stop()

            # ==========================================
            # 2. ОСНОВНОЙ РАСЧЕТ И ЛОГИКА КОЛОНОК
            # ==========================================
            for i, row_idx in enumerate(range(2, max_rows)):
                status_text.text(f"Анализ строки {row_idx} (аптека {ws.cell(row=row_idx, column=4).value})...")
                progress_bar.progress((i + 1) / total_items)
                
                visit_id = ws.cell(row=row_idx, column=1).value
                if not visit_id:
                    visit_id = f"Строка_{row_idx}"

                visit_date = ws.cell(row=row_idx, column=2).value
                pharmacy_id = ws.cell(row=row_idx, column=4).value 
                shelf_name = str(ws.cell(row=row_idx, column=8).value).strip() 
                
                # Извлекаем месяц визита
                try:
                    visit_month = pd.to_datetime(visit_date).month
                except:
                    visit_month = 1

                # Берем ссылки из 10 и 11 колонок
                cell_10 = ws.cell(row=row_idx, column=10)
                cell_11 = ws.cell(row=row_idx, column=11)
                
                url_10 = cell_10.hyperlink.target if hasattr(cell_10, 'hyperlink') and cell_10.hyperlink else str(cell_10.value)
                url_11 = cell_11.hyperlink.target if hasattr(cell_11, 'hyperlink') and cell_11.hyperlink else str(cell_11.value)
                
                image_url = None
                if url_10 and str(url_10).startswith("http"):
                    image_url = str(url_10).strip()
                elif url_11 and str(url_11).startswith("http"):
                    image_url = str(url_11).strip()

                if not image_url:
                    continue

                base_row_data = {
                    "ID Визита": visit_id,
                    "Дата": visit_date,
                    "ID Аптеки": pharmacy_id,
                    "Название полки": shelf_name,
                    "Ссылка на фото": image_url
                }

                temp_filename = f"temp_visit_{row_idx}.jpg"
                try:
                    # 1. Пытаемся скачать фото, притворяясь обычным браузером
                    req_headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)'}
                    response = requests.get(image_url, headers=req_headers, timeout=15)
                    
                    if response.status_code == 200:
                        with open(temp_filename, "wb") as f:
                            f.write(response.content)
                    else:
                        error_log.error(f"Строка {row_idx}: Oson заблокировал скачивание фото (Ошибка {response.status_code})")
                        continue
                            
                    # 2. Передаем фото в нейросеть ТОЛЬКО если оно успешно сохранилось
                    # 2. Передаем фото в нейросеть через ПРЯМОЙ REST API (без багов библиотеки)
                    if os.path.exists(temp_filename):
                        import base64
                        with open(temp_filename, "rb") as image_file:
                            img_b64 = base64.b64encode(image_file.read()).decode("ascii")
                        
                        # Прямой запрос к серверу (работает для 4 версии)
                        api_url = f"https://detect.roboflow.com/uz_ir_pharmacy/4?api_key={api_key_input}&confidence=40&overlap=30"
                        resp = requests.post(api_url, data=img_b64, headers={"Content-Type": "application/x-www-form-urlencoded"})
                        
                        if resp.status_code != 200:
                            error_log.error(f"Строка {row_idx}: Ошибка серверов Roboflow ({resp.status_code}): {resp.text}")
                            continue
                            
                        prediction = resp.json()
                        found_items = prediction.get("predictions", [])
                    else:
                        error_log.error(f"Строка {row_idx}: Файл не скачался.")
                        continue

                    # Подсчет факта с фото
                    shelf_fact = {}
                    for item in found_items:
                        cls = item["class"]
                        shelf_fact[cls] = shelf_fact.get(cls, 0) + 1

                    # Достаем все варианты планов для этого месяца и полки
                    available_options = matrix_plans.get(visit_month, {}).get(shelf_name, {})
                    
                    best_option_score = -1
                    best_option_data = {}

                    if not available_options:
                        best_option_data["Сыгравший вариант"] = "-"
                        best_option_data["ИТОГО ПО ПОЛКЕ (%)"] = "Нет Плана"
                        best_option_data["Факт полки (шт)"] = sum(shelf_fact.values())
                        best_option_data["План полки (шт)"] = "-"
                        best_option_data["Разница полки (шт)"] = "-"
                        for product in all_unique_products:
                            readable_name = PRODUCT_NAMES.get(product, product)
                            best_option_data[f"% {readable_name}"] = "Нет плана"
                            best_option_data[f"Шт. {readable_name}"] = shelf_fact.get(product, 0)
                    else:
                        # Перебор всех вариантов и поиск наилучшего
                        for option_name, current_plan in available_options.items():
                            temp_data = {}
                            sum_of_percentages = 0
                            planned_items_count = 0
                            total_plan_packs = 0
                            total_fact_packs = 0 

                            for product in all_unique_products:
                                readable_name = PRODUCT_NAMES.get(product, product)
                                col_pct = f"% {readable_name}"
                                col_fact = f"Шт. {readable_name}"

                                target = current_plan.get(product, 0)
                                fact = shelf_fact.get(product, 0)
                                
                                temp_data[col_fact] = fact

                                if target > 0:
                                    pct = min((fact / target) * 100, 100)
                                    temp_data[col_pct] = round(pct, 1)
                                    sum_of_percentages += pct
                                    planned_items_count += 1
                                    total_plan_packs += target
                                    total_fact_packs += fact
                                else:
                                    temp_data[col_pct] = "-"

                            if planned_items_count > 0:
                                overall_pct = round(sum_of_percentages / planned_items_count, 1)
                                temp_data["ИТОГО ПО ПОЛКЕ (%)"] = overall_pct
                                temp_data["Факт полки (шт)"] = total_fact_packs
                                temp_data["План полки (шт)"] = total_plan_packs
                                temp_data["Разница полки (шт)"] = total_fact_packs - total_plan_packs
                            else:
                                temp_data["ИТОГО ПО ПОЛКЕ (%)"] = 0
                                temp_data["Факт полки (шт)"] = total_fact_packs
                                temp_data["План полки (шт)"] = total_plan_packs
                                temp_data["Разница полки (шт)"] = total_fact_packs - total_plan_packs
                                
                            temp_data["Сыгравший вариант"] = option_name

                            # Если этот вариант лучше, запоминаем его
                            if overall_pct > best_option_score:
                                best_option_score = overall_pct
                                best_option_data = temp_data.copy()

                    # Собираем итоговую строку
                    final_row_data = {**base_row_data, **best_option_data}
                    master_results.append(final_row_data)

                except Exception as e:
                    # ТЕПЕРЬ ВСЕ СКРЫТЫЕ ОШИБКИ БУДУТ ВИДНЫ КРАСНЫМ ЦВЕТОМ ПРЯМО НА САЙТЕ
                    error_log.error(f"Сбой в строке {row_idx}: {e}")
                finally:
                    if os.path.exists(temp_filename):
                        os.remove(temp_filename)

            # ==========================================
            # 3. СБОРКА И КОМБИНИРОВАННАЯ РАСКРАСКА
            # ==========================================
            if master_results:
                base_cols = ["ID Визита", "Дата", "ID Аптеки", "Название полки", "Сыгравший вариант", "Ссылка на фото", 
                             "ИТОГО ПО ПОЛКЕ (%)", "Факт полки (шт)", "План полки (шт)", "Разница полки (шт)"]
                
                product_cols = []
                for p in all_unique_products:
                    readable_name = PRODUCT_NAMES.get(p, p)
                    product_cols.extend([f"% {readable_name}", f"Шт. {readable_name}"])
                
                df_final = pd.DataFrame(master_results)[base_cols + product_cols]
                
                st.success("✅ Анализ завершен!")
                st.dataframe(df_final)
                
                buffer = io.BytesIO()
                
                color_totals_header = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
                product_colors_list = [
                    PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
                    PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"),
                    PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
                    PatternFill(start_color="E4DFEC", end_color="E4DFEC", fill_type="solid"),
                    PatternFill(start_color="FDCEE8", end_color="FDCEE8", fill_type="solid"),
                    PatternFill(start_color="D0F0C0", end_color="D0F0C0", fill_type="solid")
                ]

                fill_red = PatternFill(start_color="FCE8E6", end_color="FCE8E6", fill_type="solid")   

                with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
                    df_final.to_excel(writer, index=False, sheet_name='Аудит')
                    ws_out = writer.sheets['Аудит']
                    
                    headers = [cell.value for cell in ws_out[1]]
                    
                    for col_idx, header in enumerate(headers, 1):
                        ws_out.cell(row=1, column=col_idx).font = Font(bold=True)
                        
                        if header in ["ИТОГО ПО ПОЛКЕ (%)", "Факт полки (шт)", "План полки (шт)", "Разница полки (шт)", "Сыгравший вариант"]:
                            ws_out.cell(row=1, column=col_idx).fill = color_totals_header
                        else:
                            for p_idx, p in enumerate(all_unique_products):
                                readable_name = PRODUCT_NAMES.get(p, p)
                                if header == f"% {readable_name}" or header == f"Шт. {readable_name}":
                                    color_to_use = product_colors_list[p_idx % len(product_colors_list)]
                                    ws_out.cell(row=1, column=col_idx).fill = color_to_use
                                    break
                    
                    for row_idx in range(2, len(df_final) + 2):
                        for col_idx, header in enumerate(headers, 1):
                            cell_value = ws_out.cell(row=row_idx, column=col_idx).value
                            
                            if header == "ИТОГО ПО ПОЛКЕ (%)" or str(header).startswith("%"):
                                if isinstance(cell_value, (int, float)) and cell_value < 100:
                                    ws_out.cell(row=row_idx, column=col_idx).fill = fill_red
                            elif header == "Разница полки (шт)":
                                if isinstance(cell_value, (int, float)) and cell_value < 0:
                                    ws_out.cell(row=row_idx, column=col_idx).fill = fill_red

                st.download_button(
                    label="📥 Скачать итоговый отчет",
                    data=buffer.getvalue(),
                    file_name="Итоговый_отчет.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            else:
                st.warning("Данные не найдены.")
                
        except Exception as e:
            st.error(f"Произошла ошибка: {e}")